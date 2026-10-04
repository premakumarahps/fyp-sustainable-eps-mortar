import pptx
from pptx import Presentation

PPTX_PATH = r"d:\1.Antigravity Projects\14_Final_Year_Project\Group 24.pptx"

def inspect_group_elements():
    prs = Presentation(PPTX_PATH)
    slide = prs.slides[0]
    
    def walk_shape(shape, depth=0):
        indent = "  " * depth
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                txt = "".join([r.text for r in p.runs]).strip()
                if txt:
                    print(f"{indent}[TXT] {txt}")
        if shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.GROUP:
            for child in shape.shapes:
                walk_shape(child, depth + 1)
                
    for idx, shape in enumerate(slide.shapes):
        print(f"\n=== Shape [{idx:02d}]: {shape.name} ===")
        walk_shape(shape)

if __name__ == '__main__':
    inspect_group_elements()
