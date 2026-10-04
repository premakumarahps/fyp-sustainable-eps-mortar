import pptx
from pptx import Presentation

PPTX_PATH = r"d:\1.Antigravity Projects\14_Final_Year_Project\Group 24.pptx"

def deep_walk():
    prs = Presentation(PPTX_PATH)
    slide = prs.slides[0]
    
    def walk(shapes, prefix=""):
        for i, s in enumerate(shapes):
            path = f"{prefix}/{s.name}" if prefix else s.name
            if s.has_text_frame:
                txt = s.text_frame.text.strip()
                if txt:
                    print(f"\n[{path}]")
                    for pi, p in enumerate(s.text_frame.paragraphs):
                        pt = "".join([r.text for r in p.runs]).strip()
                        if pt:
                            print(f"  P{pi:02d}: {pt}")
            if s.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.GROUP:
                walk(s.shapes, path)

    walk(slide.shapes)

if __name__ == '__main__':
    deep_walk()
