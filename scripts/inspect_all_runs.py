import pptx
from pptx import Presentation

PPTX_PATH = r"d:\1.Antigravity Projects\14_Final_Year_Project\Group 24.pptx"

def inspect_all_runs():
    prs = Presentation(PPTX_PATH)
    slide = prs.slides[0]
    
    def process_shape(shape, path=""):
        current_path = f"{path} > {shape.name}" if path else shape.name
        if shape.has_text_frame:
            tf = shape.text_frame
            full_txt = tf.text.strip()
            if full_txt:
                print(f"\n[{current_path}]")
                for p_idx, p in enumerate(tf.paragraphs):
                    p_txt = "".join([r.text for r in p.runs]).strip()
                    if p_txt:
                        print(f"  P{p_idx:02d}: {p_txt}")
        if shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.GROUP:
            for child in shape.shapes:
                process_shape(child, current_path)

    for shape in slide.shapes:
        process_shape(shape)

if __name__ == '__main__':
    inspect_all_runs()
