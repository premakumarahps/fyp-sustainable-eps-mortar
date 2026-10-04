import pptx
from pptx import Presentation
from pptx.util import Inches, Pt

PPTX_PATH = r"d:\1.Antigravity Projects\14_Final_Year_Project\Group 24.pptx"

def inspect_shapes():
    prs = Presentation(PPTX_PATH)
    slide = prs.slides[0]
    print(f"Slide width: {prs.slide_width / 914400:.2f} inches ({prs.slide_width / 360000:.1f} cm)")
    print(f"Slide height: {prs.slide_height / 914400:.2f} inches ({prs.slide_height / 360000:.1f} cm)")
    print(f"Total shapes on slide: {len(slide.shapes)}")
    
    for idx, shape in enumerate(slide.shapes):
        shape_type = shape.shape_type
        name = shape.name
        left = shape.left / 914400
        top = shape.top / 914400
        width = shape.width / 914400
        height = shape.height / 914400
        
        has_text = shape.has_text_frame
        text_snippet = ""
        if has_text:
            text_snippet = shape.text_frame.text.replace('\n', ' ')[:60]
            
        print(f"Shape [{idx:02d}] Name: {name:20s} Type: {str(shape_type):15s} Pos: ({left:.1f}, {top:.1f}) Size: ({width:.1f}x{height:.1f}) Text: {text_snippet}")

if __name__ == '__main__':
    inspect_shapes()
