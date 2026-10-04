import pptx
from pptx import Presentation
import re

SRC_PPTX = r"d:\1.Antigravity Projects\14_Final_Year_Project\Group 24.pptx"
OUT_PPTX = r"d:\1.Antigravity Projects\14_Final_Year_Project\Group 24_FINAL_CORRECTED.pptx"

# Dictionary of replacements to fix every corrupted, missing, or broken string
REPLACEMENTS = {
    # 1. Abstract
    "Conventional EPS mortar achieves weight reduction but suffers severe strength loss and brittle handling failure due to a weak, porous interface (). We engineered a tripartite upgrade combining surface-coated EPS,  pozzolanic Rice Husk Ash (RHA), and polypropylene micro-fibers. The upgraded mortar delivers  compressive strength recovery ( at  EPS),  flexural enhancement, and up to  toughness leap with zero density penalty (). SEM-EDS confirms a reduction (, dense C-S-H gel), slashing embodied carbon intensity by per .":
    "Conventional EPS mortar achieves 14–24% weight reduction but suffers severe strength loss and brittle handling failure due to a weak, porous interface (ITZ Ca/Si ≈ 24.85). We engineered a tripartite upgrade combining surface-coated EPS, 10% pozzolanic Rice Husk Ash (RHA), and 0.5% polypropylene micro-fibers. The upgraded mortar delivers +53.9% compressive strength recovery (15.62 MPa at 20% EPS), +71.0% flexural enhancement, and up to +105% toughness leap with zero density penalty (1734–1969 kg/m³). SEM-EDS confirms a 77.3% ITZ Ca/Si reduction (24.85 → 5.65, dense C-S-H gel), slashing embodied carbon intensity by 39.6% per MPa.",

    # 2. Introduction: The Paradox
    "Conventional EPS mortar cuts precast panel weight (1424% lighter) but suffers severe compressive loss and brittle failure, causing high edge breakage during demolding and transit. Raw EPS acts as a smooth, hydrophobic void that destroys the mechanical integrity and toughness needed to survive handling.":
    "Conventional EPS mortar cuts precast panel weight (14–24% lighter) but suffers severe compressive loss and brittle failure, causing high edge breakage during demolding and transit. Raw EPS acts as a smooth, hydrophobic void that destroys the mechanical integrity and toughness needed to survive handling.",

    # 3. Introduction: Microstructural Degradation Mechanisms
    "Uncoated EPS creates a porous, Portlandite-dominated Interfacial Transition Zone () - a network of pre-formed cleavage planes. Its ultra-low density also causes severe buoyancy and aggregate segregation during casting.":
    "Uncoated EPS creates a porous, Portlandite-dominated Interfacial Transition Zone (ITZ Ca/Si ≈ 24.85) - a network of pre-formed cleavage planes. Its ultra-low density also causes severe buoyancy and aggregate segregation during casting.",

    # 4. Introduction: Research Objectives
    "Engineer a tripartite ternary system (PVAc-coated EPS,  pozzolanic RHA, and PP fibers) to chemically heal the ITZ into dense C-S-H gel, mechanically lock beads, and arrest brittleness under .":
    "Engineer a tripartite ternary system (PVAc-coated EPS, 10% pozzolanic RHA, and 0.5% PP fibers) to chemically heal the ITZ into dense C-S-H gel, mechanically lock beads, and arrest brittleness under 2000 kg/m³.",

    # 5. Section C: Fiber Optimization
    "On the RHA + coated-EPS matrix, PP fiber proved optimum, boosting toughness to () and flexural capacity to . Dosages exceeding triggered fiber balling and void entrapment, inducing a steep degradation plateau ( dropped to ).":
    "On the 10% RHA + 20% coated-EPS matrix, 0.5% PP fiber proved optimum, boosting toughness to 18.24 mJ/mm³ (+24%) and flexural capacity to 3.59 MPa. Dosages exceeding 1.0% triggered fiber balling and void entrapment, inducing a steep degradation plateau (fc dropped to 11.06 MPa).",

    # 6. Table 1 Caption
    "Table 1 : The line graphs showing the steep downward trendline of strength vs. fiber dosage":
    "Table 1 : Comparative Performance Evaluation of Conventional vs. Upgraded EPS Mortar",

    # 7. Conclusions: Mechanical Recovery
    "Optimized ternary mortar ( RHA + coated EPS + PP) achieves compressive (), flexural (), and toughness at .":
    "Optimized ternary mortar (10% RHA + 20% coated EPS + 0.5% PP) achieves 15.62 MPa compressive (+53.9%), 3.59 MPa flexural (+71%), and 18.24 mJ/mm³ toughness at 1969 kg/m³.",

    # 8. Conclusions: Microstructural Proof
    "drops by (), confirming dense C-S-H gel bonding.":
    "ITZ Ca/Si drops by -77.3% (24.85 → 5.65), confirming dense C-S-H gel bonding.",

    # 9. Conclusions: Industrial Impact
    "Enhances cost efficiency by and slashes embodied carbon by per using standard precast batching.":
    "Enhances cost efficiency by +33.7% and slashes embodied carbon by 39.6% per MPa using standard precast batching.",

    # 10. Methodology & Reference encoding glitches
    "aggregatepaste boundary": "aggregate–paste boundary",
    "RHAcoated-EPS matrix": "RHA–coated-EPS matrix",
    "pp. 94269430": "pp. 9426–9430"
}

def fix_presentation():
    prs = Presentation(SRC_PPTX)
    slide = prs.slides[0]
    
    fixed_count = 0

    def clean_text_frame(tf):
        nonlocal fixed_count
        full_text = tf.text
        for old_txt, new_txt in REPLACEMENTS.items():
            if old_txt in full_text:
                # Replace in paragraphs
                for p in tf.paragraphs:
                    p_full = "".join([r.text for r in p.runs])
                    if old_txt in p_full:
                        # Clear existing runs and set new text on first run
                        if len(p.runs) > 0:
                            p.runs[0].text = new_txt
                            for r in p.runs[1:]:
                                r.text = ""
                        else:
                            p.text = new_txt
                        fixed_count += 1
                        print(f"[FIXED] Replaced snippet: {old_txt[:40]}...")
            # Also handle partial substring replacements like encoding 
            for p in tf.paragraphs:
                for r in p.runs:
                    if "" in r.text:
                        r.text = r.text.replace("", "–")
                        fixed_count += 1
                        print("[FIXED] Cleaned encoding character  -> –")

    def process_shape(shape):
        if shape.has_text_frame:
            clean_text_frame(shape.text_frame)
        if shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.GROUP:
            for child in shape.shapes:
                process_shape(child)

    for shape in slide.shapes:
        process_shape(shape)

    prs.save(OUT_PPTX)
    print(f"\nAll fixes applied! Total fixes: {fixed_count}")
    print(f"Saved corrected presentation to: {OUT_PPTX}")

if __name__ == '__main__':
    fix_presentation()
